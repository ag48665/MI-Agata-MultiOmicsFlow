from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


# =====================================================
# PATHS
# =====================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]

OUTPUTS_DIR = PROJECT_ROOT / "outputs"
REPORTS_DIR = OUTPUTS_DIR / "reports"
SUBMISSION_DIR = OUTPUTS_DIR / "submission"

MAIN_FIGURES_DIR = SUBMISSION_DIR / "main_figures"
MAIN_TABLES_DIR = SUBMISSION_DIR / "main_tables"

RESULTS_FILE = (
    REPORTS_DIR
    / "month2_manuscript_results.md"
)

DISCUSSION_FILE = (
    REPORTS_DIR
    / "month2_manuscript_discussion.md"
)

CONCLUSIONS_FILE = (
    REPORTS_DIR
    / "month2_manuscript_conclusions.md"
)

FIGURE_CAPTIONS_FILE = (
    SUBMISSION_DIR
    / "figure_captions.md"
)

TABLE_CAPTIONS_FILE = (
    SUBMISSION_DIR
    / "table_captions.md"
)

OUTPUT_FILE = (
    SUBMISSION_DIR
    / "MI_Agata_MultiOmics_Month2_Final_Manuscript.docx"
)


# =====================================================
# FIGURES
# =====================================================

FIGURES = [
    (
        "Figure 1",
        MAIN_FIGURES_DIR
        / "Figure_1_model_validation.png",
        (
            "Leakage-aware model validation. "
            "Comparison of clinical-genomic and "
            "clinical-genomic-epigenomic model performance."
        ),
    ),
    (
        "Figure 2",
        MAIN_FIGURES_DIR
        / "Figure_2_precision_recall.png",
        (
            "Precision-recall performance of the "
            "clinical-genomic-epigenomic classifier."
        ),
    ),
    (
        "Figure 3",
        MAIN_FIGURES_DIR
        / "Figure_3_calibration.png",
        (
            "Probability calibration curves comparing "
            "uncalibrated, sigmoid and isotonic predictions."
        ),
    ),
    (
        "Figure 4",
        MAIN_FIGURES_DIR
        / "Figure_4_confusion_matrix.png",
        (
            "Confusion matrix at the optimized "
            "classification threshold."
        ),
    ),
    (
        "Figure 5",
        MAIN_FIGURES_DIR
        / "Figure_5_survival.png",
        (
            "Kaplan-Meier survival analysis according "
            "to model-derived risk groups."
        ),
    ),
    (
        "Figure 6",
        MAIN_FIGURES_DIR
        / "Figure_6_feature_importance.png",
        (
            "Combined feature importance across the "
            "multi-omics modelling workflow."
        ),
    ),
]


# =====================================================
# TABLES
# =====================================================

TABLES = [
    (
        "Table 1",
        MAIN_TABLES_DIR
        / "Table_1_model_validation.csv",
        (
            "Leakage-aware cross-validation performance "
            "of the clinical-genomic and epigenomic models."
        ),
    ),
    (
        "Table 2",
        MAIN_TABLES_DIR
        / "Table_2_threshold_optimization.csv",
        (
            "Comparison of default and optimized "
            "classification thresholds."
        ),
    ),
    (
        "Table 3",
        MAIN_TABLES_DIR
        / "Table_3_calibration.csv",
        (
            "Comparison of probability calibration methods."
        ),
    ),
    (
        "Table 4",
        MAIN_TABLES_DIR
        / "Table_4_final_model_metrics.csv",
        (
            "Performance metrics of the final calibrated model."
        ),
    ),
    (
        "Table 5",
        MAIN_TABLES_DIR
        / "Table_5_key_results.csv",
        (
            "Summary of the principal Month 2 analytical results."
        ),
    ),
]


# =====================================================
# HELPERS
# =====================================================

def read_text(file_path: Path) -> str:
    """Read required text file."""

    if not file_path.exists():
        raise FileNotFoundError(
            f"Required file not found:\n{file_path}"
        )

    return file_path.read_text(
        encoding="utf-8"
    )


def clean_markdown_line(line: str) -> str:
    """Remove simple Markdown formatting."""

    line = line.strip()

    line = line.replace("**", "")
    line = line.replace("`", "")

    return line


def add_markdown_section(
    document: Document,
    text: str,
) -> None:
    """Add simple Markdown content to Word."""

    for raw_line in text.splitlines():

        line = raw_line.strip()

        if not line:
            continue

        if line.startswith("# "):

            document.add_heading(
                clean_markdown_line(
                    line[2:]
                ),
                level=1,
            )

        elif line.startswith("## "):

            document.add_heading(
                clean_markdown_line(
                    line[3:]
                ),
                level=2,
            )

        elif line.startswith("### "):

            document.add_heading(
                clean_markdown_line(
                    line[4:]
                ),
                level=3,
            )

        elif line.startswith("- "):

            paragraph = document.add_paragraph(
                style="List Bullet"
            )

            paragraph.add_run(
                clean_markdown_line(
                    line[2:]
                )
            )

        else:

            paragraph = document.add_paragraph()

            paragraph.add_run(
                clean_markdown_line(
                    line
                )
            )


def add_dataframe(
    document: Document,
    df: pd.DataFrame,
) -> None:
    """Add DataFrame as Word table."""

    df = df.copy()

    # Keep Word tables manageable.
    if len(df) > 30:
        df = df.head(30)

    table = document.add_table(
        rows=1,
        cols=len(df.columns),
    )

    table.style = "Table Grid"

    header_cells = table.rows[0].cells

    for index, column in enumerate(
        df.columns
    ):

        header_cells[index].text = str(
            column
        )

    for _, row in df.iterrows():

        cells = table.add_row().cells

        for index, value in enumerate(
            row
        ):

            if pd.isna(value):
                text = ""
            elif isinstance(
                value,
                float,
            ):
                text = f"{value:.4f}"
            else:
                text = str(value)

            cells[index].text = text


def configure_document(
    document: Document,
) -> None:
    """Set basic manuscript typography."""

    styles = document.styles

    normal = styles["Normal"]

    normal.font.name = "Arial"
    normal.font.size = Pt(10)

    for style_name in [
        "Title",
        "Heading 1",
        "Heading 2",
        "Heading 3",
    ]:

        styles[
            style_name
        ].font.name = "Arial"


# =====================================================
# MAIN
# =====================================================

def main() -> None:
    """Build final Month 2 Word manuscript."""

    print(
        "Building final Word manuscript..."
    )

    results_text = read_text(
        RESULTS_FILE
    )

    discussion_text = read_text(
        DISCUSSION_FILE
    )

    conclusions_text = read_text(
        CONCLUSIONS_FILE
    )

    document = Document()

    configure_document(
        document
    )

    # =================================================
    # TITLE PAGE
    # =================================================

    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    title_run = title.add_run(
        "Clinical-Genomic-Epigenomic "
        "Integration for Multi-Omics "
        "Predictive Modelling in Breast Cancer"
    )

    title_run.bold = True
    title_run.font.size = Pt(18)

    subtitle = document.add_paragraph()

    subtitle.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    subtitle.add_run(
        "MI-Agata-MultiOmicsFlow — Month 2"
    ).italic = True

    document.add_paragraph()

    status = document.add_paragraph()

    status.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    status.add_run(
        "Development-stage internally validated analysis"
    )

    document.add_page_break()

    # =================================================
    # ABSTRACT PLACEHOLDER
    # =================================================

    document.add_heading(
        "Abstract",
        level=1,
    )

    document.add_paragraph(
        (
            "This manuscript reports the Month 2 extension "
            "of the MI-Agata-MultiOmicsFlow pipeline, "
            "integrating DNA methylation with clinical "
            "and genomic data. The analysis includes "
            "leakage-aware cross-validation, threshold "
            "optimisation, probability calibration, "
            "survival modelling and final model validation."
        )
    )

    document.add_paragraph(
        (
            "A structured final abstract should be prepared "
            "after completion of the full Introduction and "
            "Materials and Methods sections."
        )
    )

    document.add_page_break()

    # =================================================
    # RESULTS
    # =================================================

    add_markdown_section(
        document,
        results_text,
    )

    document.add_page_break()

    # =================================================
    # MAIN TABLES
    # =================================================

    document.add_heading(
        "Main Tables",
        level=1,
    )

    for (
        label,
        file_path,
        caption,
    ) in TABLES:

        document.add_heading(
            label,
            level=2,
        )

        paragraph = document.add_paragraph()

        run = paragraph.add_run(
            caption
        )

        run.bold = True

        if not file_path.exists():

            document.add_paragraph(
                (
                    "[Table file not found: "
                    f"{file_path.name}]"
                )
            )

            continue

        df = pd.read_csv(
            file_path,
            low_memory=False,
        )

        add_dataframe(
            document,
            df,
        )

        document.add_paragraph()

    document.add_page_break()

    # =================================================
    # MAIN FIGURES
    # =================================================

    document.add_heading(
        "Main Figures",
        level=1,
    )

    for (
        label,
        file_path,
        caption,
    ) in FIGURES:

        document.add_heading(
            label,
            level=2,
        )

        if file_path.exists():

            paragraph = document.add_paragraph()

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            run = paragraph.add_run()

            run.add_picture(
                str(file_path),
                width=Inches(6.2),
            )

        else:

            document.add_paragraph(
                (
                    "[Figure file not found: "
                    f"{file_path.name}]"
                )
            )

        caption_paragraph = (
            document.add_paragraph()
        )

        caption_paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        caption_run = (
            caption_paragraph.add_run(
                f"{label}. {caption}"
            )
        )

        caption_run.italic = True

        document.add_paragraph()

    document.add_page_break()

    # =================================================
    # DISCUSSION
    # =================================================

    add_markdown_section(
        document,
        discussion_text,
    )

    document.add_page_break()

    # =================================================
    # CONCLUSIONS
    # =================================================

    add_markdown_section(
        document,
        conclusions_text,
    )

    document.add_page_break()

    # =================================================
    # LIMITATIONS
    # =================================================

    document.add_heading(
        "Study Limitations",
        level=1,
    )

    limitations = [
        (
            "All predictive performance estimates are "
            "based on internal validation within the "
            "available cohort."
        ),
        (
            "Independent external validation has not "
            "yet been performed."
        ),
        (
            "High-dimensional methylation data may "
            "increase feature-selection instability."
        ),
        (
            "The selected classification threshold "
            "is a development-stage operating point "
            "and is not a clinically validated cutoff."
        ),
        (
            "The final model and calibrated probabilities "
            "must not be used for clinical decision-making "
            "without independent validation."
        ),
    ]

    for limitation in limitations:

        document.add_paragraph(
            limitation,
            style="List Bullet",
        )

    # =================================================
    # REPRODUCIBILITY
    # =================================================

    document.add_heading(
        "Reproducibility and Supporting Materials",
        level=1,
    )

    document.add_paragraph(
        (
            "Complete analytical tables, figures, "
            "validation results, model configuration, "
            "quality-control reports and submission "
            "materials are stored in the outputs "
            "directory of MI-Agata-MultiOmicsFlow."
        )
    )

    document.add_paragraph(
        (
            "The Month 2 pipeline passed the final "
            "automated analytical QC audit and the "
            "submission-material QC audit."
        )
    )

    # =================================================
    # SAVE
    # =================================================

    document.save(
        OUTPUT_FILE
    )

    print(
        "\n========================================"
    )

    print(
        "FINAL WORD MANUSCRIPT COMPLETE"
    )

    print(
        "========================================"
    )

    print(
        "\nManuscript saved to:\n"
        f"{OUTPUT_FILE}"
    )

    print(
        "\nFinished successfully."
    )


if __name__ == "__main__":
    main()